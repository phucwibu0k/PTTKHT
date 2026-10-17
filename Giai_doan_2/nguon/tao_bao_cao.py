"""Build the phase 2 report and the phase 1 addendum using installed tools."""
from copy import deepcopy
from pathlib import Path
from tempfile import TemporaryDirectory
import re
import subprocess
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table

SRC=Path(__file__).resolve().parent
PROJECT=SRC.parents[1]
BASE=PROJECT/'Giai_doan_1/Bao_cao_Giai_doan_1.docx'

def txt(el):
    return ''.join(n.text or '' for n in el.iter(qn('w:t')))

def format_document(doc):
    for name in ['Normal','Body Text','First Paragraph','Compact','Caption']:
        if name not in doc.styles: continue
        st=doc.styles[name];st.font.name='Liberation Serif';st.font.size=Pt(11.5)
        st.paragraph_format.space_after=Pt(5);st.paragraph_format.line_spacing=1.05
    for name,size in [('Heading 1',16),('Heading 2',13),('Title',18)]:
        st=doc.styles[name];st.font.name='Liberation Sans';st.font.size=Pt(size)
        st.font.color.rgb=RGBColor.from_string('17354E')
        st.paragraph_format.space_before=Pt(6);st.paragraph_format.space_after=Pt(8)
        st.paragraph_format.keep_with_next=True
    for s in doc.sections:
        s.header_distance=Cm(.8);s.footer_distance=Cm(.8)
        s.different_first_page_header_footer=False
    parts={}
    for s in doc.sections:
        for part in [s.header,s.footer,s.first_page_header,s.first_page_footer,s.even_page_header,s.even_page_footer]:
            parts[str(part.part.partname)]=part
    for part in parts.values():
        for e in list(part._element):part._element.remove(e)
        part._element.append(OxmlElement('w:p'))
    seen=set()
    for s in doc.sections:
        name=str(s.header.part.partname)
        if name in seen:continue
        seen.add(name)
        p=s.header.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after=Pt(0)
        r=p.add_run();r.font.size=Pt(11)
        for tag,value in [('w:fldChar','begin'),('w:instrText',' PAGE '),('w:fldChar','end')]:
            el=OxmlElement(tag)
            if tag=='w:fldChar':el.set(qn('w:fldCharType'),value)
            else:el.text=value;el.set(qn('xml:space'),'preserve')
            r._r.append(el)

def page(doc,landscape=False):
    current=doc.sections[-1]
    was_landscape=current.page_width>current.page_height
    if was_landscape!=landscape:
        s=doc.add_section(WD_SECTION_START.NEW_PAGE)
        s.orientation=WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        s.page_width,s.page_height=(Cm(29.7),Cm(21)) if landscape else (Cm(21),Cm(29.7))
        s.different_first_page_header_footer=False
    else:doc.add_page_break()

def add_table(doc,el):
    source=Table(el,doc._body)
    values=[[c.text for c in row.cells] for row in source.rows]
    n=len(values[0])
    widths=([4.8,12.2] if n==2 else [4.5,6.4,6.1] if n==3 else [4.1,4.3,4.3,4.3])
    if values[0][0]=='Lớp':widths=[3.8,6.1,7.1]
    if values[0][0]=='Kiểu / ký hiệu':widths=[4.5,12.5]
    if values[0][0]=='Use Case con':widths=[6,7,4]
    table=doc.add_table(rows=len(values),cols=n);table.style='Table Grid';table.autofit=False
    for col,w in zip(table.columns,widths):col.width=Cm(w)
    for i,(row,vals) in enumerate(zip(table.rows,values)):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if i==0:pr.append(OxmlElement('w:tblHeader'))
        for cell,value,w in zip(row.cells,vals,widths):
            cell.width=Cm(w);cell.text=value
            if i==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EAF0F5');cell._tc.get_or_add_tcPr().append(sh)
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3)
                p.paragraph_format.line_spacing=1.0
                for r in p.runs:r.font.size=Pt(10.5);r.bold=(i==0)

def append_md(doc,source):
    with TemporaryDirectory(prefix='pttkht-doc-') as tmp:
        tmp=Path(tmp)
        md=source.read_text().replace('\\landscape','LANDSCAPEPAGE').replace('\\page','PORTRAITPAGE')
        (tmp/'input.md').write_text(md)
        subprocess.run(['pandoc',str(tmp/'input.md'),'-o',str(tmp/'raw.docx'),'--resource-path='+str(SRC)],check=True)
        raw=Document(tmp/'raw.docx')
        image_rels={rid:doc.part.relate_to(rel.target_part,rel.reltype) for rid,rel in raw.part.rels.items() if 'image' in rel.reltype}
        abstract={a.get(qn('w:abstractNumId')):a.find(qn('w:lvl')).find(qn('w:numFmt')).get(qn('w:val')) for a in raw.part.numbering_part.element.findall(qn('w:abstractNum'))}
        formats={n.get(qn('w:numId')):abstract[n.find(qn('w:abstractNumId')).get(qn('w:val'))] for n in raw.part.numbering_part.element.findall(qn('w:num'))}
        counts={}
        for el in raw.element.body:
            if el.tag in [qn('w:sectPr'),qn('w:bookmarkStart'),qn('w:bookmarkEnd')]:continue
            text=txt(el)
            if text in ['LANDSCAPEPAGE','PORTRAITPAGE']:
                page(doc,text=='LANDSCAPEPAGE');continue
            if re.match(r'^Hình (?:\d+|C\d+)\.',text) and el.find('.//'+qn('w:pStyle')) is not None:
                style=el.find('.//'+qn('w:pStyle')).get(qn('w:val'))
                if style=='ImageCaption':continue
            if el.tag==qn('w:tbl'):
                add_table(doc,el);continue
            new=deepcopy(el)
            for bookmark in list(new.iter(qn('w:bookmarkStart')))+list(new.iter(qn('w:bookmarkEnd'))):
                bookmark.getparent().remove(bookmark)
            if new.tag==qn('w:p'):
                pr=new.find(qn('w:pPr'));np=pr.find(qn('w:numPr')) if pr is not None else None
                if np is not None:
                    nid=np.find(qn('w:numId')).get(qn('w:val'));counts[nid]=counts.get(nid,0)+1
                    prefix=str(counts[nid])+'. ' if formats[nid]=='decimal' else '• '
                    r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=prefix;t.set(qn('xml:space'),'preserve');r.append(t);new.insert(1,r);pr.remove(np)
            for blip in new.iter(qn('a:blip')):
                rid=blip.get(qn('r:embed'))
                if rid in image_rels:blip.set(qn('r:embed'),image_rels[rid])
            for inline in new.iter(qn('wp:inline')):
                pr=new.find(qn('w:pPr'))
                if pr is None:
                    pr=OxmlElement('w:pPr');new.insert(0,pr)
                jc=pr.find(qn('w:jc'))
                if jc is None:jc=OxmlElement('w:jc');pr.append(jc)
                jc.set(qn('w:val'),'center')
                e=inline.find(qn('wp:extent'));w,h=int(e.get('cx')),int(e.get('cy'))
                landscape=doc.sections[-1].page_width>doc.sections[-1].page_height
                maxw,maxh=(Cm(25.5),Cm(14.3)) if landscape else (Cm(17),Cm(21.5))
                scale=min(maxw/w,maxh/h)
                for ext in [e,*inline.iter(qn('a:ext'))]:ext.set('cx',str(int(w*scale)));ext.set('cy',str(int(h*scale)))
            doc.element.body.insert(len(doc.element.body)-1,new)

def make_phase2():
    doc=Document(BASE)
    # Reuse the user's approved cover, including its embedded logo and placeholders.
    elements=list(doc.element.body);cut=None
    for i,el in enumerate(elements):
        if any(b.get(qn('w:type'))=='page' for b in el.iter(qn('w:br'))):cut=i;break
    assert cut is not None
    for el in elements[cut:]:
        if el.tag!=qn('w:sectPr'):doc.element.body.remove(el)
    for p in doc.paragraphs:
        for r in p.runs:
            if 'GIAI ĐOẠN 1' in r.text or 'Phân tích yêu cầu và Use Case' in r.text:
                r.text=r.text.replace('GIAI ĐOẠN 1','GIAI ĐOẠN 2').replace('Phân tích yêu cầu và Use Case','Thiết kế lớp và hành vi hệ thống')
    page(doc)
    doc.add_heading('MỤC LỤC',0)
    blocks=re.split(r'\\(?:page|landscape)\s*\n', (SRC/'noi_dung.md').read_text())
    for i,block in enumerate(blocks,3):
        title=re.search(r'^#{1,2} (.+)$',block,re.M)[1]
        p=doc.add_paragraph();p.paragraph_format.tab_stops.add_tab_stop(Cm(16.8),WD_TAB_ALIGNMENT.RIGHT)
        p.add_run(title+'\t'+str(i))
    page(doc)
    append_md(doc,SRC/'noi_dung.md')
    format_document(doc)
    doc.core_properties.title='Giai đoạn 2 – Thiết kế hệ thống quản lý phòng khám tư nhân'
    doc.core_properties.subject='Thiết kế lớp và hành vi hệ thống'
    doc.save(SRC.parent/'Bao_cao_Giai_doan_2.docx')
    return len(blocks)+2

def make_update():
    doc=Document(BASE)
    for p in doc.paragraphs:
        if p.text.startswith('Bộ tài liệu giai đoạn 1'):
            p.insert_paragraph_before('Phụ lục C. Cập nhật Use Case phân rã chi tiết ................................ 17')
            break
    page(doc)
    append_md(doc,SRC/'bo_sung_giai_doan_1.md')
    format_document(doc)
    doc.core_properties.title='Giai đoạn 1 cập nhật – bổ sung Use Case phân rã'
    doc.save(PROJECT/'Giai_doan_1/Bao_cao_Giai_doan_1_cap_nhat.docx')
    original=(PROJECT/'Giai_doan_1/nguon/noi_dung.md').read_text()
    appendix=(SRC/'bo_sung_giai_doan_1.md').read_text()
    for name in ['phan_ra_lich_hen','phan_ra_kham_benh','phan_ra_thu_phi']:
        appendix=appendix.replace(f']({name}.png)',f'](../../Giai_doan_2/nguon/{name}.png)')
    (PROJECT/'Giai_doan_1/nguon/noi_dung_cap_nhat.md').write_text(original+'\n\\page\n\n'+appendix)

if __name__=='__main__':
    expected=make_phase2();make_update()
    for folder,name in [(SRC.parent,'Bao_cao_Giai_doan_2.docx'),(PROJECT/'Giai_doan_1','Bao_cao_Giai_doan_1_cap_nhat.docx')]:
        subprocess.run(['libreoffice','-env:UserInstallation=file:///tmp/lo_pttkht_phase2','--headless','--convert-to','pdf','--outdir',str(folder),str(folder/name)],check=True)
    print(f'Phase 2 planned pages: {expected}; phase 1 update planned pages: 22. Verify rendered pagination.')
